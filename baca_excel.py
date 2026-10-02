"""Baca sheet layout rak (mis. 'Rak Existing') menjadi struktur data untuk tampak samping."""
import re
import openpyxl

KODE = re.compile(r"[A-D]\d\d")


def _baru(cell):
    """Sel berwarna (bukan putih/tanpa isi) dianggap racking baru (abu-abu di Excel)."""
    f = cell.fill
    if not f or f.fill_type != "solid":
        return 0
    g = f.fgColor
    if g.type == "theme" and g.theme == 0:
        return 0
    if g.type == "rgb" and str(g.rgb).upper() in ("FFFFFFFF", "00FFFFFF"):
        return 0
    return 1


def baca_rak(file, sheet=None):
    """file: path atau file-like. Return dict(racks=[...], pp={kode: PP di tabel sheet})."""
    wb_f = openpyxl.load_workbook(file)
    ws = wb_f[sheet] if sheet else wb_f.active
    if hasattr(file, "seek"):
        file.seek(0)
    wv = openpyxl.load_workbook(file, data_only=True)[ws.title]

    heads = [
        (c.row, c.column, c.value)
        for row in wv.iter_rows(min_row=11)
        for c in row
        if isinstance(c.value, str) and KODE.fullmatch(c.value) and c.column >= 5
    ]
    hrows = sorted({h[0] for h in heads})
    racks, covered = {}, set()
    for hr in hrows:
        hs = sorted([h for h in heads if h[0] == hr], key=lambda h: h[1])
        for i, (_, c, code) in enumerate(hs):
            nxt = [h[0] for h in heads if h[1] == c and h[0] > hr]
            end = (nxt[0] - 1) if nxt else wv.max_row
            nc = hs[i + 1][1] if i + 1 < len(hs) else c + 2
            cols = list(range(c, min(c + 2, nc)))
            cells = []
            for rr in range(hr + 1, min(end, 91) + 1):
                for cc in cols:
                    v = wv.cell(rr, cc).value
                    if isinstance(v, (int, float)) and v > 0:
                        cells.append([rr, cc, int(v), _baru(ws.cell(rr, cc))])
                        covered.add(cc)
            if cells:
                racks[code] = dict(code=code, grp=code[0], cells=cells)

    # kolom bernilai di luar blok rak (mis. deret "Samping" kolom CM)
    for cc in range(5, 100):
        if cc in covered:
            continue
        cells = [
            [rr, cc, int(wv.cell(rr, cc).value), _baru(ws.cell(rr, cc))]
            for rr in range(12, 92)
            if isinstance(wv.cell(rr, cc).value, (int, float))
        ]
        if len(cells) >= 5:
            racks["SMP"] = dict(code="SMP", grp="S", cells=cells)

    pp = {}
    for r in range(2, 9):
        for c in range(1, 40):
            v = wv.cell(r, c).value
            n = wv.cell(r, c + 1).value
            if isinstance(v, str) and KODE.fullmatch(v) and isinstance(n, (int, float)):
                pp[v] = int(n)
    return dict(racks=list(racks.values()), pp=pp)
